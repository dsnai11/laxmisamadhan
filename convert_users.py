import openpyxl
import csv
import os
import shutil

source_xlsx = r'C:/Users/sjaid/.gemini/antigravity/brain/b7686816-f2c0-4bdc-83fa-d1009a5f5811/.user_uploaded/media_1792234306164_a084648d.xlsx'
target_xlsx = r'bulk_users_converted.xlsx'
target_csv = r'bulk_users_converted.csv'

if os.path.exists(source_xlsx):
    wb_src = openpyxl.load_workbook(source_xlsx)
    ws_src = wb_src.active
    dest_headers = ['employeeId', 'firstName', 'department name', 'department raiser name', 'desginations', 'BRANCH', 'STATE', 'MOBILE NO', 'Mail Id']
    rows_data = [dest_headers]

    for r in range(2, ws_src.max_row + 1):
        code = ws_src.cell(row=r, column=1).value
        name = ws_src.cell(row=r, column=2).value
        desig = ws_src.cell(row=r, column=3).value
        dept = ws_src.cell(row=r, column=4).value
        branch = ws_src.cell(row=r, column=5).value
        state = ws_src.cell(row=r, column=6).value
        mobile = ws_src.cell(row=r, column=7).value
        mail = ws_src.cell(row=r, column=8).value
        if not code and not name and not mail:
            continue
        rows_data.append([
            str(code).strip() if code is not None else '',
            str(name).strip() if name is not None else '',
            '',
            str(dept).strip() if dept is not None else '',
            str(desig).strip() if desig is not None else '',
            str(branch).strip() if branch is not None else '',
            str(state).strip() if state is not None else '',
            str(mobile).strip() if mobile is not None else '',
            str(mail).strip() if mail is not None else ''
        ])

    with open(target_csv, 'w', newline='', encoding='utf-8') as f:
        csv.writer(f).writerows(rows_data)

    wb_dest = openpyxl.Workbook()
    ws_dest = wb_dest.active
    ws_dest.title = 'Bulk Import Users'
    for row in rows_data:
        ws_dest.append(row)
    for col in ws_dest.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_dest.column_dimensions[col_letter].width = max(max_len + 3, 12)
    wb_dest.save(target_xlsx)
    print(f'Converted {len(rows_data)-1} users.')
